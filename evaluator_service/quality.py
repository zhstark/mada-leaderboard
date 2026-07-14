from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .errors import FormatValidationError


QUALITY_HEADERS = ("序号", "糖度", "酸度")


@dataclass(frozen=True)
class QualityMeasurement:
    sample_id: int
    sugar: float
    acid: float


def load_quality_workbook(path: Path) -> dict[int, QualityMeasurement]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001 - malformed Office XML can raise parser-specific errors.
        raise FormatValidationError(f"{path.name} 不是有效的 xlsx 文件") from exc

    try:
        if len(workbook.worksheets) != 1:
            raise FormatValidationError(f"{path.name} 必须且只能包含一个工作表")

        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows()
        try:
            header_cells = next(rows)
        except StopIteration as exc:
            raise FormatValidationError(f"{path.name} 是空工作簿") from exc

        headers = tuple(cell.value for cell in header_cells[:3])
        extra_headers = [cell.value for cell in header_cells[3:] if cell.value not in (None, "")]
        if headers != QUALITY_HEADERS or extra_headers:
            expected = "、".join(QUALITY_HEADERS)
            raise FormatValidationError(f"{path.name} 表头必须严格为: {expected}")

        measurements: dict[int, QualityMeasurement] = {}
        for row_number, cells in enumerate(rows, start=2):
            values = [cell.value for cell in cells]
            if all(value in (None, "") for value in values):
                continue
            if len(values) < 3 or any(value in (None, "") for value in values[:3]):
                raise FormatValidationError(f"{path.name}:{row_number} 序号、糖度和酸度不能为空")
            if any(value not in (None, "") for value in values[3:]):
                raise FormatValidationError(f"{path.name}:{row_number} 包含表头之外的额外列")
            if any(cell.data_type == "f" for cell in cells[:3]):
                raise FormatValidationError(f"{path.name}:{row_number} 不允许使用公式")

            sample_id = _parse_sample_id(values[0], path.name, row_number)
            sugar = _parse_measurement(values[1], "糖度", path.name, row_number)
            acid = _parse_measurement(values[2], "酸度", path.name, row_number)
            if sample_id in measurements:
                raise FormatValidationError(f"{path.name}:{row_number} 序号重复: {sample_id}")
            measurements[sample_id] = QualityMeasurement(sample_id, sugar, acid)

        if not measurements:
            raise FormatValidationError(f"{path.name} 不包含有效数据行")
        return measurements
    finally:
        workbook.close()


def validate_quality_submission(prediction_path: Path, ground_truth_path: Path) -> None:
    ground_truth = load_quality_workbook(ground_truth_path)
    predictions = load_quality_workbook(prediction_path)
    _validate_sample_ids(ground_truth, predictions, prediction_path.name)


def load_quality_evaluation_data(
    prediction_path: Path,
    ground_truth_path: Path,
) -> tuple[dict[int, QualityMeasurement], dict[int, QualityMeasurement]]:
    ground_truth = load_quality_workbook(ground_truth_path)
    predictions = load_quality_workbook(prediction_path)
    _validate_sample_ids(ground_truth, predictions, prediction_path.name)
    return ground_truth, predictions


def _validate_sample_ids(
    ground_truth: dict[int, QualityMeasurement],
    predictions: dict[int, QualityMeasurement],
    prediction_name: str,
) -> None:
    expected_ids = set(ground_truth)
    submitted_ids = set(predictions)
    missing = sorted(expected_ids - submitted_ids)
    extra = sorted(submitted_ids - expected_ids)
    if missing:
        preview = ", ".join(map(str, missing[:5]))
        suffix = " ..." if len(missing) > 5 else ""
        raise FormatValidationError(f"{prediction_name} 缺少 {len(missing)} 个序号: {preview}{suffix}")
    if extra:
        preview = ", ".join(map(str, extra[:5]))
        suffix = " ..." if len(extra) > 5 else ""
        raise FormatValidationError(f"{prediction_name} 包含 {len(extra)} 个未知序号: {preview}{suffix}")


def _parse_sample_id(value: object, file_name: str, row_number: int) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise FormatValidationError(f"{file_name}:{row_number} 序号必须是正整数")
    sample_id = int(value)
    if not math.isfinite(float(value)) or float(value) != sample_id or sample_id <= 0:
        raise FormatValidationError(f"{file_name}:{row_number} 序号必须是正整数")
    return sample_id


def _parse_measurement(value: object, name: str, file_name: str, row_number: int) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise FormatValidationError(f"{file_name}:{row_number} {name}必须是数值")
    result = float(value)
    if not math.isfinite(result):
        raise FormatValidationError(f"{file_name}:{row_number} {name}必须是有限数值")
    if result < 0:
        raise FormatValidationError(f"{file_name}:{row_number} {name}不能为负数")
    return result
